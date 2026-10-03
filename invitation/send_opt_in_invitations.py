#!/usr/bin/env python3
"""Import opted-in contacts once, then send invitations from a separate SQLite DB."""

import argparse
import asyncio
import os
import re
import sqlite3
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_WORKBOOK = PROJECT_ROOT / "phones" / "all_phones.xlsx"
DEFAULT_PHONES_DB = Path(__file__).resolve().parent / "phones.sqlite3"
CONSENT = True

NAME_HEADERS = ("name", "title", "نام و نام خانوادگی", "نام کامل", "نام")
PHONE_HEADERS = ("phone", "mobile", "شماره تماس", "شماره تلفن", "موبایل")
SOURCE_HEADERS = ("source", "منبع")
PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
RETRYABLE_STATUSES = ("PENDING", "ERROR", "ERROR_RPC", "PAUSED_FLOOD_WAIT")


def clean_text(value: object) -> str:
    return " ".join(str(value or "").split())


def normalize_header(value: object) -> str:
    return clean_text(value).translate(str.maketrans({"ي": "ی", "ك": "ک"})).casefold()


def create_message(full_name: object, bot_link: str) -> str:
    name = clean_text(full_name)
    greeting = f"سلام {name} عزیز 👋" if name else "سلام 👋"
    return f"""{greeting}

«بازی‌چی» یه ربات دونفره برای وقت‌هاییه که می‌خوای با دوستت سرگرم بشی. 🎮

الان می‌تونی گل یا پوچ 🌸، دوز سه‌تایی ❌⭕، جرئت یا حقیقت 🎭 و حدس کلمه 🔤 بازی کنی؛ بازی‌های بیشتری هم به‌زودی اضافه می‌شن.

بازیکن‌های برتر پایان هر ماه جایزه می‌گیرن. 🏆

شروع بازی:
{bot_link}

اگر کانال یا گروهی داری و دنبال جذب عضوهای تازه‌ای، می‌تونی برای همکاری تبلیغاتی و اسپانسرشدن در بازی‌چی پیام بدی.

اگر تمایلی به دریافت پیام‌های بعدی نداری، فقط بنویس «لغو»."""


def normalize_country_code(value: str) -> str:
    digits = re.sub(r"\D", "", value.translate(PERSIAN_DIGITS))
    if not 1 <= len(digits) <= 4:
        raise ValueError("کد کشور معتبر نیست؛ برای ایران از +98 استفاده کن.")
    return f"+{digits}"


def normalize_phone(value: object, default_country_code: str) -> str:
    if normalize_country_code(default_country_code) != "+98":
        raise ValueError("برای این فهرست فقط کد کشور +98 پشتیبانی می‌شود")
    mobile = normalize_mobile(value)
    if mobile is None:
        raise ValueError("شمارهٔ موبایل باید ۱۱ رقم و با 09 شروع شود")
    return f"+98{mobile[1:]}"


def normalize_mobile(value: object) -> str | None:
    """Return an Iranian mobile in 09xxxxxxxxx form, or None if invalid."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        raw = str(int(value))
    else:
        raw = str(value)
    raw = re.sub(r"[\s\-().]+", "", raw.translate(PERSIAN_DIGITS).strip())
    if raw.startswith("+98"):
        raw = raw[3:]
    elif raw.startswith("0098"):
        raw = raw[4:]
    elif raw.startswith("98"):
        raw = raw[2:]
    raw = raw.removeprefix("0")
    if len(raw) != 10 or not raw.startswith("9") or not raw.isascii() or not raw.isdecimal():
        return None
    return f"0{raw}"


def mask_phone(phone: str) -> str:
    return f"{phone[:4]}***{phone[-3:]}" if len(phone) >= 8 else "***"


def timestamp() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def status_reason(status: object) -> str:
    return clean_text(status).split("|", 1)[0].strip() or "UNKNOWN"


def status_is_terminal(value: object) -> bool:
    return status_reason(value) not in RETRYABLE_STATUSES


def safe_error(error: BaseException, limit: int = 140) -> str:
    detail = clean_text(error).replace("|", "/")
    return f"{type(error).__name__}: {detail}"[:limit]


def validate_bot_link(value: str) -> str:
    link = value.strip().rstrip("/")
    parsed = urlparse(link)
    if parsed.scheme != "https" or parsed.netloc.casefold() not in {"t.me", "www.t.me"}:
        raise ValueError("لینک ربات باید به شکل https://t.me/BotUsername باشد")
    username = parsed.path.strip("/")
    if not username or "/" in username:
        raise ValueError("لینک ربات باید مستقیماً به یوزرنیم ربات اشاره کند")
    return link


def find_column(headers: dict[str, int], aliases: tuple[str, ...]) -> int | None:
    for alias in aliases:
        column = headers.get(normalize_header(alias))
        if column is not None:
            return column
    return None


def open_phones_database(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS phones (
            id INTEGER PRIMARY KEY,
            name TEXT NOT NULL DEFAULT '',
            phone TEXT NOT NULL,
            source TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'PENDING',
            is_send INTEGER NOT NULL DEFAULT 0 CHECK (is_send IN (0, 1))
        )
        """
    )
    connection.commit()
    return connection


def ensure_indexes(connection: sqlite3.Connection) -> None:
    connection.execute("CREATE INDEX IF NOT EXISTS idx_phones_source_id ON phones(source, id)")
    connection.execute("CREATE INDEX IF NOT EXISTS idx_phones_phone ON phones(phone)")
    connection.execute("CREATE INDEX IF NOT EXISTS idx_phones_status ON phones(status)")
    connection.commit()


def normalize_database(db_path: Path) -> tuple[int, int, list[tuple[str, int]]]:
    """Normalize stored numbers and remove rows that are not Iranian mobiles."""
    if not db_path.is_file():
        raise ValueError(f"دیتابیس شماره‌ها پیدا نشد: {db_path}")
    connection = sqlite3.connect(db_path)
    try:
        connection.create_function("mobile_phone", 1, normalize_mobile, deterministic=True)
        with connection:
            deleted = connection.execute(
                "DELETE FROM phones WHERE mobile_phone(phone) IS NULL"
            ).rowcount
            normalized = connection.execute(
                "UPDATE phones SET phone = mobile_phone(phone) WHERE phone != mobile_phone(phone)"
            ).rowcount
        sources = connection.execute(
            "SELECT source, COUNT(*) FROM phones GROUP BY source ORDER BY source"
        ).fetchall()
        return deleted, normalized, sources
    finally:
        connection.close()


def excel_phone(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)


def import_workbook(workbook_path: Path, db_path: Path) -> int:
    """Load every sheet in one transaction, preserving phone text and leading zeroes."""
    if not workbook_path.is_file():
        raise ValueError(f"فایل اکسل پیدا نشد: {workbook_path}")
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise RuntimeError("openpyxl نصب نیست؛ .venv/bin/pip install -e '.[outreach]'") from error

    connection = open_phones_database(db_path)
    workbook = None
    try:
        if connection.execute("SELECT 1 FROM phones LIMIT 1").fetchone():
            raise ValueError("جدول phones از قبل داده دارد؛ ورود دوباره متوقف شد تا وضعیت‌ها حفظ شوند")
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        total = 0
        batch: list[tuple[str, str, str]] = []
        with connection:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, None)
                if not header:
                    continue
                headers = {
                    normalize_header(value): index
                    for index, value in enumerate(header)
                    if normalize_header(value)
                }
                name_col = find_column(headers, NAME_HEADERS)
                phone_col = find_column(headers, PHONE_HEADERS)
                source_col = find_column(headers, SOURCE_HEADERS)
                if name_col is None or phone_col is None:
                    raise ValueError(f"ستون name یا phone در شیت {sheet.title!r} پیدا نشد")
                for row in rows:
                    phone = excel_phone(row[phone_col] if phone_col < len(row) else None)
                    if not phone:
                        continue
                    name = clean_text(row[name_col] if name_col < len(row) else None)
                    source = clean_text(
                        row[source_col]
                        if source_col is not None and source_col < len(row)
                        else None
                    ) or sheet.title
                    batch.append((name, phone, source))
                    if len(batch) >= 5000:
                        connection.executemany(
                            "INSERT INTO phones (name, phone, source) VALUES (?, ?, ?)", batch
                        )
                        total += len(batch)
                        batch.clear()
                        if total % 50000 == 0:
                            print(f"Imported {total:,} contacts...", flush=True)
            if batch:
                connection.executemany(
                    "INSERT INTO phones (name, phone, source) VALUES (?, ?, ?)", batch
                )
                total += len(batch)
        ensure_indexes(connection)
        return total
    finally:
        if workbook is not None:
            workbook.close()
        connection.close()


def next_contact(connection: sqlite3.Connection, source: str, after_id: int):
    return connection.execute(
        """
        SELECT id, name, phone FROM phones
        WHERE source = ? AND id > ? AND (
            status = 'PENDING' OR status LIKE 'ERROR | %'
            OR status LIKE 'ERROR_RPC | %' OR status LIKE 'PAUSED_FLOOD_WAIT | %'
        )
        ORDER BY id LIMIT 1
        """,
        (source, after_id),
    ).fetchone()


def set_status(
    connection: sqlite3.Connection, contact_id: int, status: str, *, is_send: bool = False
) -> None:
    connection.execute(
        "UPDATE phones SET status = ?, is_send = ? WHERE id = ?",
        (status, int(is_send), contact_id),
    )
    connection.commit()


def sent_count_today(connection: sqlite3.Connection) -> int:
    day = datetime.now().astimezone().date().isoformat()
    return int(connection.execute(
        "SELECT COUNT(*) FROM phones WHERE is_send = 1 AND status LIKE ?",
        (f"SENT | {day}%",),
    ).fetchone()[0])


def print_delivery_report(connection: sqlite3.Connection, source: str) -> None:
    rows = connection.execute(
        """
        SELECT CASE WHEN instr(status, ' | ') > 0
                    THEN substr(status, 1, instr(status, ' | ') - 1)
                    ELSE status END AS reason,
               COUNT(*)
        FROM phones WHERE source = ? GROUP BY reason ORDER BY reason
        """,
        (source,),
    ).fetchall()
    total = sum(count for _, count in rows)
    sent = next((count for reason, count in rows if reason == "SENT"), 0)
    print(f"\nگزارش سورس {source}: کل={total:,}، ارسال موفق={sent:,}")
    for reason, count in rows:
        print(f"  {reason}: {count:,}")


def load_dependencies():
    try:
        from telethon import TelegramClient, errors
        from telethon.tl.functions.contacts import ImportContactsRequest
        from telethon.tl.types import InputPhoneContact
    except ModuleNotFoundError as error:
        raise RuntimeError("Telethon نصب نیست؛ .venv/bin/pip install -e '.[outreach]'") from error
    return TelegramClient, errors, ImportContactsRequest, InputPhoneContact


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="ورود شماره‌ها و ارسال دعوت از phones.sqlite3")
    actions = parser.add_mutually_exclusive_group()
    actions.add_argument("--import-excel", action="store_true", help="ورود یک‌بارهٔ همهٔ شیت‌های اکسل")
    actions.add_argument("--normalize-phones", action="store_true", help="یکسان‌سازی شماره‌ها و حذف نامعتبرها")
    parser.add_argument("--db", type=Path, default=DEFAULT_PHONES_DB)
    parser.add_argument("--source", default=os.getenv("OUTREACH_SOURCE"))
    parser.add_argument("--count", type=int, default=None)
    parser.add_argument("--api-id", type=int, default=os.getenv("TELEGRAM_API_ID"))
    parser.add_argument("--api-hash", default=os.getenv("TELEGRAM_API_HASH"))
    parser.add_argument("--phone", default=os.getenv("OUTREACH_PHONE"))
    parser.add_argument(
        "--session", type=Path, default=Path(os.getenv("OUTREACH_SESSION", "data/outreach/account"))
    )
    parser.add_argument(
        "--bot-link", default=os.getenv("OUTREACH_BOT_LINK", "https://t.me/bazi_chi_admin")
    )
    parser.add_argument("--default-country-code", default="+98")
    parser.add_argument("--delay", type=float, default=120)
    parser.add_argument("--max-per-day", type=int, default=20)
    parser.add_argument("--dry-run", action="store_true", help="نمایش حداکثر پنج مخاطب و متن پیام")
    return parser.parse_args()


def validate_args(args: argparse.Namespace) -> None:
    if not CONSENT:
        raise ValueError("CONSENT باید صریحاً True باشد")
    if args.import_excel or getattr(args, "normalize_phones", False):
        return
    args.source = clean_text(args.source)
    if not args.source:
        raise ValueError("--source الزامی است")
    if args.count is None or args.count < 1:
        raise ValueError("--count باید عددی بزرگ‌تر از صفر باشد")
    if not args.db.is_file():
        raise ValueError(f"دیتابیس شماره‌ها پیدا نشد: {args.db}؛ ابتدا --import-excel را اجرا کن")
    args.bot_link = validate_bot_link(args.bot_link)
    normalize_country_code(args.default_country_code)
    if args.delay < 30:
        raise ValueError("فاصلهٔ ارسال باید حداقل ۳۰ ثانیه باشد")
    if not 1 <= args.max_per_day <= 100:
        raise ValueError("max-per-day باید بین ۱ و ۱۰۰ باشد")
    if not args.dry_run and (not args.api_id or not args.api_hash):
        raise ValueError("TELEGRAM_API_ID و TELEGRAM_API_HASH الزامی هستند")


async def run(args: argparse.Namespace) -> int:
    validate_args(args)
    if args.import_excel:
        total = import_workbook(INPUT_WORKBOOK, args.db)
        print(f"Imported {total:,} contacts into {args.db}; normalizing...", flush=True)
        deleted, normalized, sources = normalize_database(args.db)
        print(f"Deleted {deleted:,} invalid contacts; normalized {normalized:,} numbers")
        for source, count in sources:
            print(f"  {source}: {count:,}")
        return 0
    if getattr(args, "normalize_phones", False):
        deleted, normalized, sources = normalize_database(args.db)
        print(f"Deleted {deleted:,} invalid contacts; normalized {normalized:,} numbers")
        for source, count in sources:
            print(f"  {source}: {count:,}")
        return 0

    connection = sqlite3.connect(args.db)
    try:
        if args.dry_run:
            print(f"DB: {args.db}; source={args.source}")
            last_id = 0
            for _ in range(min(args.count, 5)):
                contact = next_contact(connection, args.source, last_id)
                if contact is None:
                    break
                contact_id, name, raw_phone = contact
                last_id = contact_id
                print(f"\n--- id={contact_id} / {mask_phone(raw_phone)} ---")
                print(create_message(name, args.bot_link))
            print_delivery_report(connection, args.source)
            return 0

        first_contact = next_contact(connection, args.source, 0)
        if first_contact is None:
            print("هیچ مخاطب تازه‌ای برای این سورس باقی نمانده است.")
            print_delivery_report(connection, args.source)
            return 0
        sent_today = sent_count_today(connection)
        if sent_today >= args.max_per_day:
            print(f"سقف امروز پر شده است: {sent_today}/{args.max_per_day}")
            print_delivery_report(connection, args.source)
            return 0

        TelegramClient, errors, ImportContactsRequest, InputPhoneContact = load_dependencies()
        args.session.parent.mkdir(parents=True, exist_ok=True)
        client = TelegramClient(
            str(args.session), int(args.api_id), args.api_hash,
            flood_sleep_threshold=0, request_retries=1,
        )
        try:
            if args.phone:
                await client.start(phone=args.phone)
            else:
                await client.start()
            me = await client.get_me()
            print(f"Telegram connected: id={me.id}; source={args.source}; count={args.count}")
            consecutive_errors = 0
            contact = first_contact
            attempted = 0
            last_id = 0
            requests_made = False
            while contact is not None and attempted < args.count:
                if sent_today >= args.max_per_day:
                    print(f"سقف روزانه پر شد: {sent_today}/{args.max_per_day}")
                    break
                contact_id, name, raw_phone = contact
                last_id = contact_id
                attempted += 1
                try:
                    phone = normalize_phone(raw_phone, args.default_country_code)
                except ValueError as error:
                    set_status(connection, contact_id, f"INVALID_PHONE | {safe_error(error)} | {timestamp()}")
                    contact = next_contact(connection, args.source, last_id)
                    continue

                # The original number remains in the table; avoid a second send for
                # an exact matching number in another row or source.
                duplicate = connection.execute(
                    "SELECT 1 FROM phones WHERE phone = ? AND id <> ? AND is_send = 1 LIMIT 1",
                    (raw_phone, contact_id),
                ).fetchone()
                if duplicate:
                    set_status(connection, contact_id, f"DUPLICATE_SENT | {timestamp()}")
                    contact = next_contact(connection, args.source, last_id)
                    continue

                if requests_made:
                    print(f"[WAIT] {args.delay:.0f}s")
                    await asyncio.sleep(args.delay)
                print(f"[CHECK] id={contact_id} phone={mask_phone(phone)}")
                set_status(connection, contact_id, f"SENDING | {timestamp()}")
                requests_made = True
                should_stop = False
                try:
                    result = await client(
                        ImportContactsRequest([
                            InputPhoneContact(
                                client_id=contact_id, phone=phone,
                                first_name=name or "مخاطب", last_name="",
                            )
                        ])
                    )
                    if not result.users:
                        set_status(connection, contact_id, f"NOT_ON_TELEGRAM | {timestamp()}")
                        print(f"[NOT ON TELEGRAM] id={contact_id}")
                    else:
                        user = result.users[0]
                        await client.send_message(
                            user, create_message(name, args.bot_link),
                            link_preview=False, parse_mode=None,
                        )
                        set_status(connection, contact_id, f"SENT | {timestamp()}", is_send=True)
                        sent_today += 1
                        print(f"[SENT] id={contact_id}; today={sent_today}/{args.max_per_day}")
                    consecutive_errors = 0
                except errors.FloodWaitError as error:
                    seconds = int(getattr(error, "seconds", 0))
                    set_status(connection, contact_id, f"PAUSED_FLOOD_WAIT | retry_after={seconds}s | {timestamp()}")
                    print(f"[STOP] Telegram requested FloodWait({seconds}s).")
                    should_stop = True
                except errors.PeerFloodError as error:
                    set_status(connection, contact_id, f"STOPPED_SPAM_RESTRICTION | {safe_error(error)} | {timestamp()}")
                    print("[STOP] Telegram reported an account spam restriction.")
                    should_stop = True
                except errors.UserPrivacyRestrictedError as error:
                    set_status(connection, contact_id, f"NOT_SENT_PRIVACY | {safe_error(error)} | {timestamp()}")
                    consecutive_errors = 0
                except errors.RPCError as error:
                    set_status(connection, contact_id, f"ERROR_RPC | {safe_error(error)} | {timestamp()}")
                    consecutive_errors += 1
                    print(f"[RPC ERROR] id={contact_id}: {safe_error(error)}")
                except Exception as error:  # noqa: BLE001 - persist unexpected delivery failures
                    set_status(connection, contact_id, f"ERROR | {safe_error(error)} | {timestamp()}")
                    consecutive_errors += 1
                    print(f"[ERROR] id={contact_id}: {safe_error(error)}")

                if should_stop or consecutive_errors >= 3:
                    if consecutive_errors >= 3:
                        print("[STOP] Three consecutive errors; inspect phones.sqlite3.")
                    break
                contact = next_contact(connection, args.source, last_id)
        finally:
            await client.disconnect()

        print_delivery_report(connection, args.source)
        return 0
    finally:
        connection.close()


def main() -> int:
    try:
        return asyncio.run(run(parse_args()))
    except (ValueError, RuntimeError) as error:
        print(f"خطا: {error}", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("\nمتوقف شد؛ نتیجه‌های ثبت‌شده در phones.sqlite3 مانده‌اند.", file=sys.stderr)
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
