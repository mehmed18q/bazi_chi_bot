#!/usr/bin/env python3
"""Send invitations to the project's fully opted-in, deduplicated contact workbook.

The input workbook is intentionally fixed. It is read in streaming mode and never
modified. Per-contact delivery state is kept in SQLite so a multi-million-row XLSX
does not have to be loaded and rewritten after every Telegram request.
"""

import argparse
import asyncio
import os
import re
import sqlite3
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator
from urllib.parse import urlparse


PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_WORKBOOK = PROJECT_ROOT / "phones" / "all_phones_deduplicated.xlsx"
DEFAULT_STATE_DB = PROJECT_ROOT / "data" / "outreach" / "invitations.sqlite3"

# The owner of this list has explicitly confirmed that every contact opted in.
CONSENT = True

NAME_HEADERS = ("name", "title", "نام و نام خانوادگی", "نام کامل", "نام")
PHONE_HEADERS = ("phone", "mobile", "شماره تماس", "شماره تلفن", "موبایل")
SOURCE_HEADERS = ("source", "منبع")

TERMINAL_STATUS_PREFIXES = (
    "SENT",
    "NOT_ON_TELEGRAM",
    "NOT_SENT_PRIVACY",
    "INVALID_PHONE",
    "STOPPED_SPAM_RESTRICTION",
    "DO_NOT_CONTACT",
    "UNSUBSCRIBED",
    "لغو",
)
PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


@dataclass(frozen=True, slots=True)
class ExcelContact:
    sequence: int
    sheet_name: str
    row_number: int
    name: str
    raw_phone: object
    source: str


@dataclass(frozen=True, slots=True)
class PreparedContact:
    sequence: int
    sheet_name: str
    row_number: int
    name: str
    phone: str
    source: str

    @property
    def record_key(self) -> str:
        return self.phone


def clean_text(value: object) -> str:
    return " ".join(str(value or "").split())


def normalize_header(value: object) -> str:
    return clean_text(value).translate(str.maketrans({"ي": "ی", "ك": "ک"})).casefold()


def create_message(full_name: object, bot_link: str) -> str:
    """Build plain text; an empty Excel name produces a natural generic greeting."""
    name = clean_text(full_name)
    greeting = f"سلام {name} عزیز 👋" if name else "سلام 👋"
    return f"""{greeting}

«بازی‌چی» یه ربات دونفره برای وقت‌هاییه که می‌خوای با دوستت سرگرم بشی. 🎮

الان می‌تونی گل یا پوچ 🌸، دوز سه‌تایی ❌⭕، جرئت یا حقیقت 🎭 و حدس کلمه 🔤 بازی کنی؛ بازی‌های بیشتری هم به‌زودی اضافه می‌شن.

بازیکن‌های برتر پایان هر ماه جایزه می‌گیرن. 🏆

شروع بازی:
{bot_link}

اگر کانال یا گروهی داری و دنبال جذب عضوهای تازه‌ای، می‌تونی برای همکاری تبلیغاتی و اسپانسرشدن در بازی‌چی پیام بدی.

اگر تمایلی به دریافت پیام‌های بعدی نداری، فقط بنویس «لغو»."""


def normalize_country_code(value: str) -> str:
    digits = re.sub(r"\D", "", value.translate(PERSIAN_DIGITS))
    if not 1 <= len(digits) <= 4:
        raise ValueError("کد کشور معتبر نیست؛ برای ایران از +98 استفاده کن.")
    return f"+{digits}"


def normalize_phone(value: object, default_country_code: str) -> str:
    if value is None or isinstance(value, bool):
        raise ValueError("شماره خالی است")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError("شماره در اکسل به‌صورت عدد اعشاری ذخیره شده است")
        raw = str(int(value))
    else:
        raw = str(value)
    raw = raw.translate(PERSIAN_DIGITS).strip()
    raw = re.sub(r"[\s\-()]+", "", raw)
    if raw.startswith("00"):
        raw = f"+{raw[2:]}"

    country_code = normalize_country_code(default_country_code)
    country_digits = country_code[1:]
    if raw.startswith("+"):
        normalized = f"+{re.sub(r'\D', '', raw[1:])}"
    else:
        digits = re.sub(r"\D", "", raw)
        if digits.startswith(country_digits):
            normalized = f"+{digits}"
        elif digits.startswith("0"):
            normalized = f"{country_code}{digits[1:]}"
        elif country_code == "+98" and len(digits) == 10 and digits.startswith("9"):
            normalized = f"+98{digits}"
        else:
            raise ValueError("شماره باید با کد کشور، مثل +98912...، ثبت شده باشد")

    digits = normalized[1:]
    if not 8 <= len(digits) <= 15 or digits.startswith("0"):
        raise ValueError("طول یا قالب شماره معتبر نیست")
    return normalized


def mask_phone(phone: str) -> str:
    return f"{phone[:4]}***{phone[-3:]}" if len(phone) >= 8 else "***"


def timestamp() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def status_is_terminal(value: object) -> bool:
    status = clean_text(value).upper()
    return any(status.startswith(prefix.upper()) for prefix in TERMINAL_STATUS_PREFIXES)


def validate_bot_link(value: str) -> str:
    link = value.strip().rstrip("/")
    parsed = urlparse(link)
    if parsed.scheme != "https" or parsed.netloc.casefold() not in {"t.me", "www.t.me"}:
        raise ValueError("لینک ربات باید به شکل https://t.me/BotUsername باشد")
    username = parsed.path.strip("/")
    if not username or "/" in username:
        raise ValueError("لینک ربات باید مستقیماً به یوزرنیم ربات اشاره کند")
    return link


def header_map(header_row: tuple[object, ...]) -> dict[str, int]:
    return {
        normalize_header(value): index
        for index, value in enumerate(header_row)
        if normalize_header(value)
    }


def find_column(headers: dict[str, int], aliases: tuple[str, ...]) -> int | None:
    for alias in aliases:
        column = headers.get(normalize_header(alias))
        if column is not None:
            return column
    return None


def safe_error(error: BaseException, limit: int = 140) -> str:
    text = clean_text(error).replace("|", "/")
    return f"{type(error).__name__}: {text}"[:limit]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "ارسال دعوت بازی‌چی به فهرست ثابت all_phones_deduplicated.xlsx؛ "
            "همهٔ مخاطبان این فایل رضایت داده‌اند"
        )
    )
    parser.add_argument(
        "--api-id", type=int, default=os.getenv("TELEGRAM_API_ID"), help="API ID اکانت تلگرام"
    )
    parser.add_argument(
        "--api-hash", default=os.getenv("TELEGRAM_API_HASH"), help="API hash اکانت تلگرام"
    )
    parser.add_argument(
        "--phone",
        default=os.getenv("OUTREACH_PHONE"),
        help="شمارهٔ اکانت با کد کشور؛ اگر خالی باشد Telethon تعاملی می‌پرسد",
    )
    parser.add_argument(
        "--session",
        type=Path,
        default=Path(os.getenv("OUTREACH_SESSION", "data/outreach/account")),
        help="مسیر session؛ اطلاعات ورود بعد از اولین اجرا در آن می‌ماند",
    )
    parser.add_argument(
        "--state-db",
        type=Path,
        default=Path(os.getenv("OUTREACH_STATE_DB", str(DEFAULT_STATE_DB))),
        help="SQLite وضعیت ارسال، لغو و خطاها",
    )
    parser.add_argument(
        "--bot-link", default=os.getenv("OUTREACH_BOT_LINK"), help="مثل https://t.me/MyBot"
    )
    parser.add_argument("--default-country-code", default="+98")
    parser.add_argument(
        "--delay",
        type=float,
        default=120,
        help="فاصلهٔ ثابت بین بررسی مخاطبان، بر حسب ثانیه؛ حداقل مجاز ۳۰ است",
    )
    parser.add_argument(
        "--max-per-day",
        type=int,
        default=20,
        help="سقف محافظه‌کارانهٔ ارسال موفق در هر روز",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="فقط پنج مخاطب آماده و متن پیام را بررسی کن؛ به تلگرام وصل نشو",
    )
    return parser.parse_args()


def validate_args(args: argparse.Namespace) -> None:
    if not CONSENT:
        raise ValueError("CONSENT باید برای این فهرست صریحاً True باشد")
    if not INPUT_WORKBOOK.is_file() or INPUT_WORKBOOK.suffix.casefold() != ".xlsx":
        raise ValueError(f"فایل ورودی ثابت پیدا نشد: {INPUT_WORKBOOK}")
    if not args.bot_link:
        raise ValueError("--bot-link یا OUTREACH_BOT_LINK الزامی است")
    args.bot_link = validate_bot_link(args.bot_link)
    normalize_country_code(args.default_country_code)
    if args.delay < 30:
        raise ValueError("فاصلهٔ ارسال باید حداقل ۳۰ ثانیه باشد")
    if not 1 <= args.max_per_day <= 100:
        raise ValueError("max-per-day باید بین ۱ و ۱۰۰ باشد")
    if not args.dry_run and (not args.api_id or not args.api_hash):
        raise ValueError("TELEGRAM_API_ID و TELEGRAM_API_HASH یا آرگومان معادل آن‌ها الزامی است")


def load_dependencies() -> tuple[Any, Any, Any, Any]:
    try:
        from openpyxl import load_workbook
        from telethon import TelegramClient, errors
        from telethon.tl.functions.contacts import ImportContactsRequest
        from telethon.tl.types import InputPhoneContact
    except ModuleNotFoundError as error:
        raise RuntimeError(
            "وابستگی‌ها نصب نیستند؛ اجرا کن: .venv/bin/pip install -e '.[outreach]'"
        ) from error
    return load_workbook, TelegramClient, errors, (ImportContactsRequest, InputPhoneContact)


def open_state_database(path: Path | str) -> sqlite3.Connection:
    if path != ":memory:":
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path)
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS invitation_status (
            record_key TEXT PRIMARY KEY,
            phone TEXT NOT NULL,
            name TEXT NOT NULL,
            source TEXT NOT NULL,
            sheet_name TEXT NOT NULL,
            row_number INTEGER NOT NULL,
            status TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            sent_at TEXT,
            telegram_id INTEGER
        )
        """
    )
    connection.execute(
        "CREATE INDEX IF NOT EXISTS idx_invitation_status_sent_at ON invitation_status(sent_at)"
    )
    connection.commit()
    return connection


def stored_status(connection: sqlite3.Connection, record_key: str) -> str:
    row = connection.execute(
        "SELECT status FROM invitation_status WHERE record_key = ?", (record_key,)
    ).fetchone()
    return row[0] if row else ""


def store_status(
    connection: sqlite3.Connection,
    contact: ExcelContact | PreparedContact,
    record_key: str,
    phone: str,
    status: str,
    *,
    sent_at: str | None = None,
    telegram_id: int | None = None,
) -> None:
    now = timestamp()
    connection.execute(
        """
        INSERT INTO invitation_status (
            record_key, phone, name, source, sheet_name, row_number,
            status, updated_at, sent_at, telegram_id
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(record_key) DO UPDATE SET
            phone = excluded.phone,
            name = excluded.name,
            source = excluded.source,
            sheet_name = excluded.sheet_name,
            row_number = excluded.row_number,
            status = excluded.status,
            updated_at = excluded.updated_at,
            sent_at = COALESCE(excluded.sent_at, invitation_status.sent_at),
            telegram_id = COALESCE(excluded.telegram_id, invitation_status.telegram_id)
        """,
        (
            record_key,
            phone,
            contact.name,
            contact.source,
            contact.sheet_name,
            contact.row_number,
            status,
            now,
            sent_at,
            telegram_id,
        ),
    )
    connection.commit()


def sent_count_today(connection: sqlite3.Connection, today: str) -> int:
    return int(
        connection.execute(
            "SELECT COUNT(*) FROM invitation_status WHERE substr(sent_at, 1, 10) = ?",
            (today,),
        ).fetchone()[0]
    )


def iter_excel_contacts(load_workbook: Any) -> Iterator[ExcelContact]:
    workbook = load_workbook(INPUT_WORKBOOK, read_only=True, data_only=True)
    sequence = 0
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            headers = header_map(tuple(header))
            name_column = find_column(headers, NAME_HEADERS)
            phone_column = find_column(headers, PHONE_HEADERS)
            source_column = find_column(headers, SOURCE_HEADERS)
            if name_column is None:
                raise ValueError(f"ستون name در شیت {worksheet.title!r} پیدا نشد")
            if phone_column is None:
                raise ValueError(f"ستون phone در شیت {worksheet.title!r} پیدا نشد")

            for row_number, row in enumerate(rows, start=2):
                raw_phone = row[phone_column] if phone_column < len(row) else None
                if raw_phone in (None, ""):
                    continue
                sequence += 1
                name = clean_text(row[name_column] if name_column < len(row) else None)
                source = clean_text(
                    row[source_column]
                    if source_column is not None and source_column < len(row)
                    else ""
                )
                yield ExcelContact(
                    sequence=sequence,
                    sheet_name=worksheet.title,
                    row_number=row_number,
                    name=name,
                    raw_phone=raw_phone,
                    source=source,
                )
    finally:
        workbook.close()


def iter_ready_contacts(
    contacts: Iterator[ExcelContact],
    connection: sqlite3.Connection,
    default_country_code: str,
    *,
    persist_invalid: bool,
) -> Iterator[PreparedContact]:
    for contact in contacts:
        try:
            phone = normalize_phone(contact.raw_phone, default_country_code)
        except ValueError as error:
            record_key = f"invalid:{contact.sheet_name}:{contact.row_number}"
            if status_is_terminal(stored_status(connection, record_key)):
                continue
            if persist_invalid:
                store_status(
                    connection,
                    contact,
                    record_key,
                    "",
                    f"INVALID_PHONE | {safe_error(error)} | {timestamp()}",
                )
            continue

        if status_is_terminal(stored_status(connection, phone)):
            continue
        yield PreparedContact(
            sequence=contact.sequence,
            sheet_name=contact.sheet_name,
            row_number=contact.row_number,
            name=contact.name,
            phone=phone,
            source=contact.source,
        )


async def run(args: argparse.Namespace) -> int:
    validate_args(args)
    load_workbook, TelegramClient, errors, contact_types = load_dependencies()
    ImportContactsRequest, InputPhoneContact = contact_types

    state_path: Path | str = ":memory:" if args.dry_run else args.state_db
    state = open_state_database(state_path)
    contacts = iter_excel_contacts(load_workbook)
    ready = iter_ready_contacts(
        contacts,
        state,
        args.default_country_code,
        persist_invalid=not args.dry_run,
    )

    if args.dry_run:
        try:
            preview = []
            for contact in ready:
                preview.append(contact)
                if len(preview) >= 5:
                    break
            print(f"Input: {INPUT_WORKBOOK}")
            print(f"CONSENT={CONSENT}; preview_ready={len(preview)}")
            for contact in preview:
                print(
                    f"\n--- {contact.sheet_name} row {contact.row_number} / "
                    f"{mask_phone(contact.phone)} / source={contact.source or '-'} ---"
                )
                print(create_message(contact.name, args.bot_link))
            return 0
        finally:
            ready.close()
            contacts.close()
            state.close()

    today = datetime.now().astimezone().date().isoformat()
    sent_count = sent_count_today(state, today)
    if sent_count >= args.max_per_day:
        print(f"سقف امروز قبلاً پر شده است: {sent_count}/{args.max_per_day}")
        ready.close()
        contacts.close()
        state.close()
        return 0

    first_contact = next(ready, None)
    if first_contact is None:
        print("هیچ مخاطب تازه‌ای برای بررسی باقی نمانده است.")
        ready.close()
        contacts.close()
        state.close()
        return 0

    args.session.parent.mkdir(parents=True, exist_ok=True)
    client = TelegramClient(
        str(args.session),
        int(args.api_id),
        args.api_hash,
        flood_sleep_threshold=0,
        request_retries=1,
    )
    consecutive_errors = 0
    current: PreparedContact | None = first_contact
    try:
        if args.phone:
            await client.start(phone=args.phone)
        else:
            await client.start()
        me = await client.get_me()
        print(
            f"Telegram connected: id={me.id}; input={INPUT_WORKBOOK.name}; "
            f"CONSENT={CONSENT}; sent_today={sent_count}"
        )

        while current is not None:
            if sent_count >= args.max_per_day:
                print(f"سقف روزانه پر شد: {sent_count}/{args.max_per_day}")
                break
            print(
                f"[CHECK] sheet={current.sheet_name} row={current.row_number} "
                f"phone={mask_phone(current.phone)}"
            )
            should_stop = False
            try:
                result = await client(
                    ImportContactsRequest(
                        [
                            InputPhoneContact(
                                client_id=current.sequence,
                                phone=current.phone,
                                first_name=current.name or "مخاطب",
                                last_name="",
                            )
                        ]
                    )
                )
                if not result.users:
                    store_status(
                        state,
                        current,
                        current.record_key,
                        current.phone,
                        f"NOT_ON_TELEGRAM | {timestamp()}",
                    )
                    print(f"[NOT ON TELEGRAM] {current.sheet_name} row={current.row_number}")
                    consecutive_errors = 0
                else:
                    user = result.users[0]
                    await client.send_message(
                        user,
                        create_message(current.name, args.bot_link),
                        link_preview=False,
                        parse_mode=None,
                    )
                    sent_at = timestamp()
                    store_status(
                        state,
                        current,
                        current.record_key,
                        current.phone,
                        f"SENT | {sent_at} | telegram_id={user.id}",
                        sent_at=sent_at,
                        telegram_id=user.id,
                    )
                    sent_count += 1
                    consecutive_errors = 0
                    print(f"[SENT] today={sent_count}/{args.max_per_day}")
            except errors.FloodWaitError as error:
                seconds = int(getattr(error, "seconds", 0))
                store_status(
                    state,
                    current,
                    current.record_key,
                    current.phone,
                    f"PAUSED_FLOOD_WAIT | retry_after={seconds}s | {timestamp()}",
                )
                print(f"[STOP] Telegram requested FloodWait({seconds}s); no automatic retry.")
                should_stop = True
            except errors.PeerFloodError as error:
                store_status(
                    state,
                    current,
                    current.record_key,
                    current.phone,
                    f"STOPPED_SPAM_RESTRICTION | {safe_error(error)} | {timestamp()}",
                )
                print("[STOP] Telegram reported an account spam restriction.")
                should_stop = True
            except errors.UserPrivacyRestrictedError as error:
                store_status(
                    state,
                    current,
                    current.record_key,
                    current.phone,
                    f"NOT_SENT_PRIVACY | {safe_error(error)} | {timestamp()}",
                )
                consecutive_errors = 0
                print(f"[PRIVACY] {current.sheet_name} row={current.row_number}")
            except errors.RPCError as error:
                store_status(
                    state,
                    current,
                    current.record_key,
                    current.phone,
                    f"ERROR_RPC | {safe_error(error)} | {timestamp()}",
                )
                consecutive_errors += 1
                print(f"[RPC ERROR] {safe_error(error)}")
            except Exception as error:
                store_status(
                    state,
                    current,
                    current.record_key,
                    current.phone,
                    f"ERROR | {safe_error(error)} | {timestamp()}",
                )
                consecutive_errors += 1
                print(f"[ERROR] {safe_error(error)}")

            if should_stop:
                break
            if consecutive_errors >= 3:
                print("[STOP] Three consecutive errors; inspect the state database.")
                break
            if sent_count >= args.max_per_day:
                print(f"سقف روزانه پر شد: {sent_count}/{args.max_per_day}")
                break

            next_contact = next(ready, None)
            if next_contact is None:
                break
            print(f"[WAIT] {args.delay:.0f}s")
            await asyncio.sleep(args.delay)
            current = next_contact
    finally:
        await client.disconnect()
        ready.close()
        contacts.close()
        state.close()

    print(f"State database: {args.state_db}; sent_today={sent_count}/{args.max_per_day}")
    return 0


def main() -> int:
    try:
        return asyncio.run(run(parse_args()))
    except (ValueError, RuntimeError) as error:
        print(f"خطا: {error}", file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        print("\nمتوقف شد؛ نتیجه‌های قبلی در دیتابیس وضعیت ذخیره شده‌اند.", file=sys.stderr)
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
