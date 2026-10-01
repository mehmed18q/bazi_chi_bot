import importlib.util
import sqlite3
from pathlib import Path
from types import SimpleNamespace
from typing import ClassVar
from unittest.mock import AsyncMock

import pytest
from openpyxl import Workbook


def load_script():
    path = Path(__file__).parents[1] / "scripts/send_opt_in_invitations.py"
    spec = importlib.util.spec_from_file_location("send_opt_in_invitations", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sample_workbook(path):
    workbook = Workbook()
    first = workbook.active
    first.title = "vendor_user"
    first.append(["name", "phone", "source"])
    first.append(["محمد رضایی", "09121234567", "vendor_user"])
    first.append([None, "+989121234568", "vendor_user"])
    second = workbook.create_sheet("contacts_vcf")
    second.append(["name", "phone", "source"])
    second.append(["مخاطب سوم", "09121234569", "contacts_vcf"])
    test = workbook.create_sheet("test")
    test.append(["name", "phone"])
    test.append(["صادق", "09217074647"])
    workbook.save(path)


def args_for(db, tmp_path, *, source="vendor_user", count=1, dry_run=False):
    return SimpleNamespace(
        import_excel=False,
        db=db,
        source=source,
        count=count,
        api_id=12345,
        api_hash="hash",
        phone="+989121111111",
        session=tmp_path / "session/account",
        bot_link="https://t.me/GameBot",
        default_country_code="+98",
        normalize_phones=False,
        delay=30,
        max_per_day=20,
        dry_run=dry_run,
    )


def test_message_uses_single_full_name_column_and_has_generic_fallback():
    script = load_script()
    named = script.create_message("  محمد   رضایی  ", "https://t.me/GameBot")
    anonymous = script.create_message(None, "https://t.me/GameBot")
    assert named.startswith("سلام محمد رضایی عزیز 👋")
    assert anonymous.startswith("سلام 👋")
    assert "گل یا پوچ" in named and "حدس کلمه" in named
    assert "بازی‌های بیشتری" in named and "پایان هر ماه جایزه" in named
    assert "اسپانسرشدن" in named and "https://t.me/GameBot" in named


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("0912 123 4567", "+989121234567"),
        (9121234567, "+989121234567"),
        ("۰۰۹۸۹۱۲۱۲۳۴۵۶۷", "+989121234567"),
        ("+98-912-123-4567", "+989121234567"),
    ],
)
def test_phone_normalization(raw, expected):
    assert load_script().normalize_phone(raw, "+98") == expected


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("09121234567", "09121234567"),
        ("9121234567", "09121234567"),
        ("+98 912 123 4567", "09121234567"),
        ("00989121234567", "09121234567"),
        ("۹۱۲۱۲۳۴۵۶۷", "09121234567"),
        ("02532912272", None),
        ("0912123456", None),
        ("091212345678", None),
        ("9abc123456", None),
    ],
)
def test_mobile_cleanup(raw, expected):
    assert load_script().normalize_mobile(raw) == expected


def test_import_puts_numbers_and_results_in_one_table(tmp_path):
    script = load_script()
    workbook_path = tmp_path / "all_phones.xlsx"
    db = tmp_path / "phones.sqlite3"
    sample_workbook(workbook_path)

    assert script.import_workbook(workbook_path, db) == 4
    with sqlite3.connect(db) as connection:
        columns = [row[1] for row in connection.execute("PRAGMA table_info(phones)")]
        rows = connection.execute(
            "SELECT name, phone, source, status, is_send FROM phones ORDER BY id"
        ).fetchall()
    assert columns == ["id", "name", "phone", "source", "status", "is_send"]
    assert rows[0] == ("محمد رضایی", "09121234567", "vendor_user", "PENDING", 0)
    assert rows[1] == ("", "+989121234568", "vendor_user", "PENDING", 0)
    assert rows[-1] == ("صادق", "09217074647", "test", "PENDING", 0)
    with pytest.raises(ValueError, match="از قبل داده دارد"):
        script.import_workbook(workbook_path, db)


def test_normalize_database_removes_invalid_rows_and_preserves_status(tmp_path):
    script = load_script()
    db = tmp_path / "phones.sqlite3"
    with script.open_phones_database(db) as connection:
        connection.executemany(
            "INSERT INTO phones (name, phone, source, status, is_send) VALUES (?, ?, ?, ?, ?)",
            [
                ("صادق", "09217074647", "test", "SENT | 2026-10-01", 1),
                ("موبایل", "+989121234567", "vendor_user", "PENDING", 0),
                ("بدون صفر", "9121234568", "vendor_user", "PENDING", 0),
                ("ثابت", "02532912272", "contacts_vcf", "PENDING", 0),
                ("کوتاه", "0912123456", "contacts_vcf", "PENDING", 0),
            ],
        )
    deleted, normalized, sources = script.normalize_database(db)
    assert (deleted, normalized) == (2, 2)
    assert sources == [("test", 1), ("vendor_user", 2)]
    with sqlite3.connect(db) as connection:
        rows = connection.execute(
            "SELECT phone, status, is_send FROM phones ORDER BY id"
        ).fetchall()
    assert rows == [
        ("09217074647", "SENT | 2026-10-01", 1),
        ("09121234567", "PENDING", 0),
        ("09121234568", "PENDING", 0),
    ]
    assert script.normalize_database(db)[:2] == (0, 0)


async def test_import_command_normalizes_and_filters_before_use(tmp_path, monkeypatch):
    script = load_script()
    workbook_path = tmp_path / "all_phones.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "vendor_user"
    sheet.append(["name", "phone", "source"])
    sheet.append(["موبایل", "+989121234567", "vendor_user"])
    sheet.append(["ثابت", "02532912272", "vendor_user"])
    workbook.save(workbook_path)
    monkeypatch.setattr(script, "INPUT_WORKBOOK", workbook_path)
    db = tmp_path / "phones.sqlite3"

    assert await script.run(SimpleNamespace(import_excel=True, normalize_phones=False, db=db)) == 0
    with sqlite3.connect(db) as connection:
        assert connection.execute(
            "SELECT name, phone, source, status, is_send FROM phones"
        ).fetchall() == [("موبایل", "09121234567", "vendor_user", "PENDING", 0)]


async def test_send_selects_source_and_count_then_updates_same_rows(tmp_path, monkeypatch):
    script = load_script()
    workbook_path = tmp_path / "all_phones.xlsx"
    db = tmp_path / "phones.sqlite3"
    sample_workbook(workbook_path)
    script.import_workbook(workbook_path, db)

    class FakeClient:
        messages: ClassVar[list] = []
        instances: ClassVar[list] = []

        def __init__(self, *args, **kwargs):
            self.instances.append(self)

        async def start(self, **kwargs):
            return self

        async def get_me(self):
            return SimpleNamespace(id=99)

        async def __call__(self, request):
            phone = request.contacts[0].phone
            return SimpleNamespace(users=[SimpleNamespace(id=int(phone[-4:]))])

        async def send_message(self, user, message, **kwargs):
            self.messages.append((user, message, kwargs))

        async def disconnect(self):
            return None

    from telethon import errors
    from telethon.tl.functions.contacts import ImportContactsRequest
    from telethon.tl.types import InputPhoneContact

    monkeypatch.setattr(
        script,
        "load_dependencies",
        lambda: (FakeClient, errors, ImportContactsRequest, InputPhoneContact),
    )
    monkeypatch.setattr(script.asyncio, "sleep", AsyncMock())

    args = args_for(db, tmp_path)
    assert await script.run(args) == 0
    assert len(FakeClient.messages) == 1
    with sqlite3.connect(db) as connection:
        rows = connection.execute(
            "SELECT status, is_send FROM phones WHERE source = 'vendor_user' ORDER BY id"
        ).fetchall()
    assert rows[0][0].startswith("SENT | ") and rows[0][1] == 1
    assert rows[1] == ("PENDING", 0)

    assert await script.run(args) == 0
    assert len(FakeClient.messages) == 2
    assert FakeClient.messages[1][1].startswith("سلام 👋")
    assert await script.run(args) == 0
    assert len(FakeClient.instances) == 2

    test_args = args_for(db, tmp_path, source="test")
    assert await script.run(test_args) == 0
    assert FakeClient.messages[-1][1].startswith("سلام صادق عزیز 👋")
    with sqlite3.connect(db) as connection:
        assert connection.execute(
            "SELECT phone, status, is_send FROM phones WHERE source = 'test'"
        ).fetchone()[0] == "09217074647"
        assert connection.execute(
            "SELECT COUNT(*) FROM phones WHERE status = 'PENDING'"
        ).fetchone()[0] == 1
        connection.execute(
            "INSERT INTO phones (name, phone, source) VALUES (?, ?, ?)",
            ("تکراری", "09217074647", "duplicate_source"),
        )

    assert await script.run(args_for(db, tmp_path, source="duplicate_source")) == 0
    assert len(FakeClient.messages) == 3
    with sqlite3.connect(db) as connection:
        assert connection.execute(
            "SELECT status, is_send FROM phones WHERE source = 'duplicate_source'"
        ).fetchone()[0].startswith("DUPLICATE_SENT | ")


async def test_dry_run_does_not_change_status(tmp_path, monkeypatch):
    script = load_script()
    workbook_path = tmp_path / "all_phones.xlsx"
    db = tmp_path / "phones.sqlite3"
    sample_workbook(workbook_path)
    script.import_workbook(workbook_path, db)
    monkeypatch.setattr(script, "load_dependencies", lambda: pytest.fail("connected to Telegram"))

    assert await script.run(args_for(db, tmp_path, source="test", dry_run=True)) == 0
    with sqlite3.connect(db) as connection:
        assert connection.execute(
            "SELECT status, is_send FROM phones WHERE source = 'test'"
        ).fetchone() == ("PENDING", 0)
